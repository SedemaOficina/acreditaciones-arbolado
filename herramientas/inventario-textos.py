"""INVENTARIO DE TEXTOS PARA REVISIÓN (2026-10-15, Liber)

Junta en una sola hoja todo texto que ve la persona solicitante (y, desde el 24/10/2026, el paso a paso de «Grupos» del
personal; desde el 25/10/2026, también «Grupos por asignar y confirmar», de Capacitación), con la pantalla, la sección, el tipo, la condición
en que aparece y dónde se cambia:
  · docs/textos/estaticos.json (lo saca herramientas/inventario-textos.mjs de las páginas, sin ejecutar scripts);
  · los textos que pintan los scripts del sitio según el perfil, el estado o la respuesta (lista curada abajo);
  · los mensajes de la base al guardar o enviar (apd.error en db/apd/03-apd-reglas.sql);
  · los mensajes de validación del Estándar de formularios (formkit/js/validaciones.config.js).
Escribe docs/textos/inventario-textos.json y docs/textos/inventario-textos.xlsx.

  CHROMIUM=... node herramientas/inventario-textos.mjs && python3 herramientas/inventario-textos.py
"""
import json, re
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
SAL = RAIZ / 'docs' / 'textos'

PERSONA = 'Persona solicitante'
PERSONAL = 'Personal de la Secretaría'

# ── 1. Textos de las páginas ────────────────────────────────────────────────────────────────────────────────
estaticos = json.loads((SAL / 'estaticos.json').read_text(encoding='utf8'))

def corta_condicion(c):
    # De una cadena de condiciones anidadas basta la más específica («Si … es …») más las de bloque
    partes = [p.strip() for p in c.split(';')]
    si = [p for p in partes if p.startswith('Si «')]
    otras = [p for p in partes if not p.startswith('Si «')]
    return '; '.join(otras + si[-1:]) or 'Siempre'

def donde_pagina(r):
    if r.get('origen') == 'Catálogo de requisitos (base)':
        return 'Catálogo de requisitos de la base (db/apd/04-apd-catalogos.sql)'
    if r['archivo'] in ('cuenta.html', 'datos.html') and r['seccion'].endswith('Tu adscripción'):
        return 'Fragmento de adscripción (herramientas/plantillas/adscripcion.html)'
    return f"Página (sitio/{r['archivo']})" if r['archivo'] != '*.html' else 'Todas las páginas'

filas = []
for r in estaticos:
    filas.append({
        'pantalla': r['pantalla'], 'seccion': r['seccion'] or '—', 'tipo': r['tipo'],
        'cuando': corta_condicion(r['cuando']), 'texto': r['texto'], 'quien': PERSONAL if r['archivo'] in ('grupos.html', 'capacitacion.html') else PERSONA, 'donde': donde_pagina(r), 'obs': ''
    })

# La adscripción es el mismo fragmento en «Crea tu cuenta» y en «Mis datos»: una sola fila por texto
vistos = {}
unidas = []
for f in filas:
    if f['donde'].startswith('Fragmento de adscripción'):
        k = (f['tipo'], f['texto'], f['cuando'])
        if k in vistos:
            vistos[k]['pantalla'] = 'Crea tu cuenta y Mis datos'
            continue
        vistos[k] = f
    unidas.append(f)
filas = unidas

# ── 2. Textos que pintan los scripts (curados: condición en palabras) ───────────────────────────────────────
D = lambda archivo: f'Código del sitio (sitio/app/{archivo})'
dinamicos = [
    # Estados de la solicitud (apd.js · ESTADOS): etiqueta del estado y «qué sigue»
    ('Mis trámites', 'Mis solicitudes · estado', 'Estado', 'Solicitud en borrador', 'Borrador', D('apd.js')),
    ('Mis trámites', 'Mis solicitudes · qué sigue', 'Texto', 'Solicitud en borrador', 'Termina de cargar tus documentos y envía tu solicitud.', D('apd.js')),
    ('Mis trámites', 'Mis solicitudes · estado', 'Estado', 'Solicitud enviada', 'En revisión documental', D('apd.js')),
    ('Mis trámites', 'Mis solicitudes · qué sigue', 'Texto', 'Solicitud enviada', 'La Secretaría hace la revisión documental de tu solicitud: revisa que tus documentos estén completos y correctos.', D('apd.js')),
    ('Mis trámites', 'Mis solicitudes · estado', 'Estado', 'Solicitud con observaciones', 'Con observaciones', D('apd.js')),
    ('Mis trámites', 'Mis solicitudes · qué sigue', 'Texto', 'Solicitud con observaciones', 'Corrige lo que se indica abajo y vuelve a enviar tu solicitud.', D('apd.js')),
    ('Mis trámites', 'Mis solicitudes · estado', 'Estado', 'Solicitud validada en línea (con cita)', 'Cita asignada', D('apd.js')),
    ('Mis trámites', 'Mis solicitudes · qué sigue', 'Texto', 'Solicitud validada en línea (con cita)', 'Acude a tu cita con los mismos documentos que cargaste, impresos, y con tu comprobante de pago.', D('apd.js')),
    ('Mis trámites', 'Mis solicitudes · estado', 'Estado', 'Expediente completo, sin grupo', 'En espera de grupo', D('apd.js')),
    ('Mis trámites', 'Mis solicitudes · qué sigue', 'Texto', 'Expediente completo, sin grupo', 'Tu cotejo terminó. El siguiente paso es que la Secretaría te asigne un grupo del curso; te avisará las fechas y la sede.', D('apd.js')),
    ('Mis trámites', 'Mis solicitudes · estado', 'Estado', 'Inscrita en un grupo', 'Con grupo asignado', D('apd.js')),
    ('Mis trámites', 'Mis solicitudes · qué sigue', 'Texto', 'Inscrita en un grupo, sin fecha de inicio', 'Asiste a tu curso en las fechas indicadas.', D('apd.js')),
    ('Mis trámites', 'Mis solicitudes · estado', 'Estado', 'Examen aprobado', 'Aprobada', D('apd.js')),
    ('Mis trámites', 'Mis solicitudes · estado', 'Estado', 'Acreditación emitida', 'Acreditación emitida', D('apd.js')),
    ('Mis trámites', 'Mis solicitudes · estado', 'Estado', 'Examen no aprobado', 'No acreditada', D('apd.js')),
    ('Mis trámites', 'Mis solicitudes · estado', 'Estado', 'Sin la asistencia mínima', 'Sin derecho a examen', D('apd.js')),
    ('Mis trámites', 'Mis solicitudes · estado', 'Estado', 'Cancelada por la persona o la Secretaría', 'Cancelada', D('apd.js')),
    ('Mis trámites', 'Mis solicitudes · estado', 'Estado', 'Observaciones no atendidas en el plazo', 'Concluida sin atender', D('apd.js')),
    ('Mis trámites', 'Mis solicitudes · qué sigue', 'Texto', 'Observaciones no atendidas en el plazo', 'Las observaciones no se atendieron a tiempo. Puedes iniciar una solicitud nueva.', D('apd.js')),
    # Ficha de la solicitud (tramites.js)
    ('Mis trámites', 'Mis solicitudes · ficha', 'Título de ficha', 'Cada solicitud', '«Poda y derribo» o «Dictaminador técnico» · «solicitud nueva» o «renovación»', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · ficha', 'Etiqueta de dato', 'Cada solicitud', 'Folio de la solicitud', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · ficha', 'Etiqueta de dato', 'Cada solicitud', 'Fecha en que enviaste la solicitud', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · ficha', 'Etiqueta de dato', 'Solicitud con cita', 'Cita de cotejo', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · ficha', 'Etiqueta de dato', 'Desde el primer cotejo', 'Número de expediente', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · datos para tu pago', 'Título', 'Solicitud con cita', 'Datos para tu pago', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · datos para tu pago', 'Etiqueta de dato', 'Solicitud con cita', 'Concepto', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · datos para tu pago', 'Etiqueta de dato', 'Solicitud con cita', 'Monto', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · datos para tu pago', 'Etiqueta de dato', 'Solicitud con cita', 'Referencia', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · datos para tu pago', 'Nota en amarillo', 'Mientras no se confirme el concepto', 'Por confirmar: texto oficial del concepto de pago.', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · datos para tu pago', 'Nota en amarillo', 'Mientras no se confirme el monto', 'Por confirmar: lo fija la Gaceta Oficial.', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · datos para tu pago', 'Nota en amarillo', 'Mientras no se defina la referencia', 'Por confirmar: formato de la referencia de pago.', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · datos para tu pago', 'Valor', 'Si no hay monto capturado', 'Sin monto', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · datos para tu pago', 'Aviso pendiente', 'Mientras no se definan banco y forma de pago', 'Pendiente: banco, cuenta y forma de pago.', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · datos para tu pago', 'Texto', 'Solicitud con cita', 'Paga antes de tu cita y lleva el comprobante impreso. En la cita la Secretaría verifica tu pago y emite tu factura.', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · datos para tu pago', 'Texto', 'Solicitud con cita', 'Lleva también tu identificación oficial: con ella se confirma que la fotografía de tu cuenta es tuya. Si no se ve bien, ahí te toman otra.', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · curso', 'Etiqueta de dato', 'Inscrita con fecha de inicio', 'Preséntate el', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · curso', 'Nota en amarillo', 'Inscrita con fecha de inicio', 'Por confirmar: hora de entrada.', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · curso', 'Etiqueta de dato', 'Inscrita con fecha de inicio', 'El curso termina el', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · curso', 'Etiqueta de dato', 'Inscrita con fecha de inicio', 'Sede', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · curso', 'Etiqueta de dato', 'Inscrita', 'Grupo', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · qué sigue', 'Texto', 'Inscrita con fecha de inicio', 'Tu curso dura N días; necesitas asistir al menos M para presentar el examen.', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · qué sigue', 'Texto', 'Inscrita, grupo sin fecha', 'Ya tienes grupo, pero todavía no está confirmado. La Secretaría te avisará la fecha y la sede cuando se confirme.', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · observaciones', 'Aviso', 'Solicitud con observaciones', 'Lo que debes corregir: (y la lista: «Documento: observación»)', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · resultado', 'Etiqueta de dato', 'Aprobada, no acreditada o sin derecho a examen', 'Asistencia', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · resultado', 'Valor', 'Aprobada, no acreditada o sin derecho a examen', 'N de M días (mínimo X)', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · resultado', 'Etiqueta de dato', 'Aprobada o no acreditada', 'Calificación', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · resultado', 'Valor', 'Aprobada o no acreditada', '8.2 (mínima 7.0)', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · acreditación emitida', 'Etiqueta de dato', 'Acreditación emitida', 'Número de acreditación', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · acreditación emitida', 'Etiqueta de dato', 'Acreditación emitida', 'Inicio de vigencia', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · acreditación emitida', 'Etiqueta de dato', 'Acreditación emitida', 'Fin de vigencia', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · qué sigue', 'Texto', 'Acreditación emitida y entregada', 'La recogiste el DD/MM/AAAA.', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · qué sigue', 'Texto', 'Acreditación emitida, sin entregar', 'Tu acreditación está lista desde el DD/MM/AAAA. Recógela en persona con tu identificación oficial. Ahí te toman la fotografía que lleva tu credencial y, al recibirla, firmas el comprobante de recepción.', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · qué sigue', 'Aviso pendiente', 'Acreditación emitida, sin entregar, sin lugar definido', 'Pendiente: lugar, días y horario de entrega.', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · qué sigue', 'Texto', 'Examen aprobado', 'Aprobaste el examen. La Secretaría emite tu acreditación; aquí verás cuándo puedes recogerla.', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · qué sigue', 'Texto', 'Examen no aprobado', 'No alcanzaste la calificación mínima. Para obtener la acreditación haces un trámite nuevo, con su pago.', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · qué sigue', 'Texto', 'Sin la asistencia mínima', 'No alcanzaste la asistencia mínima. Para obtener la acreditación haces un trámite nuevo, con su pago.', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · ficha', 'Etiqueta de dato', 'Cancelada o concluida sin atender', 'Motivo', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · ficha', 'Botón o enlace', 'Solicitud en borrador', 'Continuar solicitud', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · ficha', 'Botón o enlace', 'Solicitud con observaciones', 'Corregir y reenviar', D('tramites.js')),
    # Oportunidades para corregir (ajustes del equipo, 25/10/2026)
    ('Mis trámites', 'Mis solicitudes · oportunidades', 'Texto', 'Con observaciones de la revisión en línea, primera vez', 'Te quedan 2 oportunidades para corregir.', D('apd.js')),
    ('Mis trámites', 'Mis solicitudes · oportunidades', 'Texto', 'Con observaciones, la última oportunidad', 'Esta es tu última oportunidad para corregir. Si en la revisión siguen las observaciones, tu solicitud se cierra y tienes que iniciar una nueva.', D('apd.js')),
    ('Mis trámites', 'Mis solicitudes · oportunidades', 'Texto', 'Con observaciones del cotejo', 'Corregir lo que no coincidió en tu cita de cotejo no cuenta entre tus oportunidades para corregir.', D('apd.js')),
    ('Solicitud', 'Corregir documentos · oportunidades', 'Texto', 'Al corregir observaciones', 'Te quedan 2 oportunidades para corregir. / Esta es tu última oportunidad para corregir. Si en la revisión siguen las observaciones, tu solicitud se cierra y tienes que iniciar una nueva.', D('solicitud.js')),
    ('Mis trámites', 'Mis solicitudes · estado', 'Estado', 'Cerrada por agotar las oportunidades de corrección', 'Cerrada: se agotaron tus oportunidades para corregir', D('apd.js')),
    ('Mis trámites', 'Mis solicitudes · qué sigue', 'Texto', 'Cerrada por agotar las oportunidades de corrección', 'En la revisión de tu última corrección siguieron las observaciones y tu solicitud se cerró. Puedes iniciar una solicitud nueva de inmediato.', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · ficha', 'Etiqueta de dato', 'Cerrada por agotar las oportunidades de corrección', 'Motivo del cierre · Se agotaron las oportunidades de corrección', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes · ficha', 'Botón o enlace', 'Cerrada por agotar las oportunidades de corrección', 'Iniciar una solicitud nueva', D('tramites.js')),
    ('Mis trámites', 'Mis solicitudes', 'Texto', 'Si no tiene solicitudes', 'Todavía no tienes solicitudes.', D('tramites.js')),
    ('Mis trámites', 'Mis trámites', 'Introducción', 'Al cargar la pantalla', 'Aquí ves tus acreditaciones y en qué va cada solicitud.', D('tramites.js')),
    ('Mis trámites', 'Mis trámites', 'Error', 'Si no se pudo cargar la información', 'No se pudo cargar tu información. Intenta de nuevo más tarde.', D('tramites.js')),
    # Mis acreditaciones
    ('Mis trámites', 'Mis acreditaciones · vigencia', 'Estado', 'Acreditación vigente', 'Vigente', D('apd.js')),
    ('Mis trámites', 'Mis acreditaciones · vigencia', 'Estado', 'Faltan 30 días o menos para que venza', 'Vigente · por vencer', D('apd.js')),
    ('Mis trámites', 'Mis acreditaciones · vigencia', 'Estado', 'Vencida (renovable o no)', 'No vigente', D('apd.js')),
    ('Mis trámites', 'Mis acreditaciones · vigencia', 'Estado', 'Emitida con vigencia futura', 'Todavía no vigente', D('apd.js')),
    ('Mis trámites', 'Mis acreditaciones · ficha', 'Etiqueta de dato', 'Cada acreditación', 'Emitida el', D('tramites.js')),
    ('Mis trámites', 'Mis acreditaciones · ficha', 'Etiqueta de dato', 'Cada acreditación', 'Vence el', D('tramites.js')),
    ('Mis trámites', 'Mis acreditaciones · ficha', 'Etiqueta de dato', 'Acreditación del registro de Kobo', 'Origen', D('tramites.js')),
    ('Mis trámites', 'Mis acreditaciones · ficha', 'Valor', 'Acreditación del registro de Kobo', 'Registro anterior de la Secretaría', D('tramites.js')),
    ('Mis trámites', 'Mis acreditaciones · ficha', 'Texto', 'Por vencer, ya renovable', 'Ya puedes renovarla. Hazlo antes de que venza para no quedarte sin acreditación vigente.', D('tramites.js')),
    ('Mis trámites', 'Mis acreditaciones · ficha', 'Texto', 'Vencida, todavía renovable', 'Ya no está vigente: no puedes trabajar con ella. Puedes renovarla hasta el DD/MM/AAAA.', D('tramites.js')),
    ('Mis trámites', 'Mis acreditaciones · ficha', 'Botón o enlace', 'Renovable', 'Renovar', D('tramites.js')),
    ('Mis trámites', 'Mis acreditaciones · ficha', 'Texto', 'Vencida hace más de un año', 'Venció hace más de un año: para obtenerla de nuevo haces una solicitud nueva.', D('tramites.js')),
    ('Mis trámites', 'Mis acreditaciones · ficha', 'Texto', 'Vigente, todavía no renovable', 'Podrás renovarla a partir del DD/MM/AAAA.', D('apd.js')),
    ('Mis trámites', 'Mis acreditaciones', 'Texto', 'Si no tiene acreditaciones', 'Todavía no tienes acreditaciones.', D('tramites.js')),
    # Cancelación (P120) y reposición de la credencial (P140), sesión de trabajo del 23/10/2026
    ('Mis trámites', 'Mis acreditaciones · vigencia', 'Estado', 'Cancelada por la Secretaría', 'Cancelada', D('apd.js')),
    ('Mis trámites', 'Mis acreditaciones · cancelada', 'Etiqueta de dato', 'Cancelada', 'Cancelada el · Motivo · Fundamento', D('tramites.js')),
    ('Mis trámites', 'Mis acreditaciones · cancelada', 'Texto', 'Cancelada', 'La Secretaría canceló tu acreditación: ya no puedes trabajar con ella y no se puede renovar.', D('tramites.js')),
    ('Mis trámites', 'Mis acreditaciones · reposición', 'Botón o enlace', 'Vigente y con la credencial ya entregada', 'Pedir la reposición de tu credencial', D('tramites.js')),
    ('Mis trámites', 'Mis acreditaciones · reposición', 'Estado', 'Reposición pedida', 'Reposición: acta en revisión · Reposición: corrige tu acta · Reposición: recoge tu credencial · Credencial repuesta', D('apd.js')),
    ('Mis trámites', 'Mis acreditaciones · reposición', 'Texto', 'Acta en revisión', 'La Secretaría revisa tu acta. Si algo falta, te lo indica aquí; si está completa, te dice cuánto pagar y dónde recoger tu credencial.', D('tramites.js')),
    ('Mis trámites', 'Mis acreditaciones · reposición', 'Botón o enlace', 'Acta con observaciones', 'Cargar el acta corregida', D('tramites.js')),
    ('Mis trámites', 'Mis acreditaciones · reposición', 'Pasos', 'Acta aceptada', 'Paga la cuota de reposición: $…, e imprime tu comprobante. · Acude a la oficina del Vivero Nezahualcóyotl con tu identificación oficial y tu comprobante de pago. · Ahí la Secretaría valida tu factura y te entrega tu credencial, con el mismo número y la misma vigencia.', D('tramites.js')),
    ('Mis trámites', 'Mis acreditaciones · reposición', 'Aviso pendiente', 'Acta aceptada', 'Pendiente: días y horario para recoger la credencial repuesta.', D('tramites.js')),
    ('Reposición de tu credencial', 'Resultado', 'Título', 'Al pedirla o al cargar el acta corregida', 'Recibimos tu solicitud de reposición · Recibimos tu acta corregida', D('reposicion.js')),
    ('Reposición de tu credencial', 'Resultado', 'Pasos', 'Al pedirla', 'La Secretaría revisa tu acta. · Te indica en «Mis trámites» si debes corregir algo o, si está completa, cuánto pagar. · Pagas la cuota de reposición y recoges tu credencial en la oficina del Vivero Nezahualcóyotl, con el mismo número y la misma vigencia.', D('reposicion.js')),
    ('Reposición de tu credencial', 'Acreditación', 'Error', 'Credencial del sistema que todavía no se entrega', 'Tu credencial todavía no se te entrega: recógela en la oficina del Vivero. La reposición es para una credencial que ya recibiste y perdiste o te robaron.', D('reposicion.js')),
    ('Verificar una acreditación', 'Resultado', 'Estado', 'Acreditación cancelada', 'Cancelada · Cancelada el DD/MM/AAAA · La Secretaría canceló esta acreditación: con ella no se puede trabajar en el arbolado urbano.', D('config-verificar.js')),
    # Sesión
    ('Pantallas con cuenta', 'Cabecera de sesión', 'Botón o enlace', 'Con sesión iniciada', 'Mis trámites', D('apd.js')),
    ('Pantallas con cuenta', 'Cabecera de sesión', 'Botón o enlace', 'Con sesión iniciada', 'Mis datos', D('apd.js')),
    ('Pantallas con cuenta', 'Cabecera de sesión', 'Botón o enlace', 'Con sesión iniciada', 'Salir', D('apd.js')),
    # Solicitud: qué puede pedir (apd.js · opciones) y textos del formulario (solicitud.js)
    ('Solicitud', 'Trámite · acreditación deshabilitada', 'Motivo', 'Ya tiene un trámite de ese tipo en curso', 'Ya tienes un trámite en curso (folio APD-S-…).', D('apd.js')),
    ('Solicitud', 'Trámite · acreditación deshabilitada', 'Motivo', 'Vigente y todavía no renovable', 'Tu acreditación PD-… vence el DD/MM/AAAA. Podrás renovarla a partir del DD/MM/AAAA.', D('apd.js')),
    ('Solicitud', 'Trámite · acreditación deshabilitada', 'Motivo', 'Emitida con vigencia futura', 'Tu acreditación PD-… todavía no entra en vigor.', D('apd.js')),
    ('Solicitud', 'Trámite · acreditación deshabilitada', 'Motivo', 'Dictaminador sin poda vigente', 'Necesitas tener vigente tu acreditación de poda y derribo.', D('apd.js')),
    ('Solicitud', 'Trámite · modalidad', 'Texto', 'Le corresponde renovación', 'Te corresponde la renovación de tu acreditación PD-…. Curso de N días.', D('solicitud.js')),
    ('Solicitud', 'Trámite · modalidad', 'Texto', 'Le corresponde solicitud nueva', 'Te corresponde una solicitud nueva: es tu primera vez o tu acreditación anterior venció hace más de un año. Curso de N días.', D('solicitud.js')),
    ('Solicitud', 'Trámite · adscripción', 'Valor', 'Persona independiente', 'Independiente: no trabajas para ninguna empresa, institución o dependencia.', D('apd.js')),
    ('Solicitud', 'Trámite · adscripción', 'Valor', 'Persona con adscripción', '«Tipo»: «Nombre de la institución».  (tipos: Empresa, consultoría o particular · Dependencia del Gobierno de la Ciudad de México · Alcaldía · Dependencia del Gobierno Federal · Institución educativa)', D('apd.js')),
    ('Solicitud', 'Trámite · adscripción', 'Valor', 'Cuenta sin adscripción', 'sin indicar. Antes de enviar, indica en «Mis datos» si trabajas para alguna empresa, institución o dependencia.', D('solicitud.js')),
    ('Solicitud', 'Trámite · experiencia', 'Ayuda (cambia)', 'Dictaminador con perfil que pide 2 años', 'Con tu perfil se piden al menos dos años.', D('solicitud.js')),
    ('Solicitud', 'Trámite · experiencia', 'Ayuda (cambia)', 'Dictaminador con perfil que pide 5 años o que podría no cumplir', 'Con tu perfil se piden al menos cinco años.', D('solicitud.js')),
    ('Solicitud', 'Trámite · experiencia', 'Ayuda (cambia)', 'Dictaminador con carrera por clasificar', 'Se piden al menos cinco años: la Secretaría revisará tu carrera cuando revise tus documentos.', D('solicitud.js')),
    ('Solicitud', 'Continuar una solicitud', 'Texto', 'Al continuar un borrador o corregir observaciones', 'Continúas la solicitud APD-S-…. La acreditación ya no cambia.', D('solicitud.js')),
    ('Solicitud', 'Continuar una solicitud', 'Error', 'Si el folio no es suyo o ya no se puede modificar', 'No encontramos la solicitud APD-S-… entre tus trámites, o ya no se puede modificar.', D('solicitud.js')),
    ('Solicitud', 'Documentos · cargado antes', 'Ayuda (cambia)', 'Al corregir: documento ya aceptado', 'Ya lo cargaste y fue aceptado; no tienes que volver a subirlo.', D('solicitud.js')),
    ('Solicitud', 'Documentos · cargado antes', 'Ayuda (cambia)', 'Al continuar un borrador: documento ya cargado', 'Ya lo cargaste; súbelo de nuevo sólo si quieres reemplazarlo.', D('solicitud.js')),
    ('Solicitud', 'Revisión · documentos', 'Valor', 'En la revisión, bajo cada documento', 'Archivo que subiste: «nombre original».pdf', D('solicitud.js')),
    ('Solicitud', 'Revisión · documentos', 'Valor', 'En la revisión, con dos o más documentos', 'N documentos', D('solicitud.js')),
    ('Solicitud', 'Solicitud', 'Error', 'Si no se pudo cargar la información', 'No se pudo cargar tu información. Intenta de nuevo más tarde.', D('solicitud.js')),
    # Inicio (requisitos: los arma inicio.js con el catálogo)
    ('Inicio', 'Requisitos', 'Título desplegable', 'Uno por acreditación y modalidad', '«Poda y derribo» o «Dictaminador técnico» · «solicitud nueva» o «renovación»', D('inicio.js')),
    ('Inicio', 'Requisitos', 'Texto', 'Requisito que se entrega en la cita', 'Se entrega en la cita de cotejo.', D('inicio.js')),
    ('Inicio', 'Requisitos', 'Texto', 'Requisito que se hace al entregar la credencial', 'Se hace al recoger tu credencial.', D('inicio.js')),
    ('Inicio', 'Requisitos', 'Texto', 'Bajo cada lista de requisitos', 'Curso de N días; asistencia mínima de M días; calificación mínima del examen: X. Pago de derechos: $… (o «pendiente de la Gaceta Oficial»).', D('inicio.js')),
    ('Inicio', 'Requisitos', 'Error', 'Si no se pudo cargar el catálogo', 'No se pudieron cargar los requisitos. Intenta de nuevo más tarde.', D('inicio.js')),
    # Cuenta y entrar (mensajes que devuelve el servidor)
    ('Crea tu cuenta', 'Al crear la cuenta', 'Error', 'Si ya hay una cuenta con esa CURP', 'Ya hay una cuenta con esta CURP. Entra con ella; si no recuerdas tu contraseña, pide ayuda a la Secretaría.', D('config-cuenta.js')),
    ('Crea tu cuenta', 'Al crear la cuenta', 'Error', 'CURP del registro de Kobo con otro nombre', 'Tu nombre no coincide con el registro anterior de la Secretaría para esta CURP. Acude a la Secretaría para aclararlo.', D('config-cuenta.js')),
    ('Crea tu cuenta', 'CURP', 'Error', 'Si la CURP no corresponde al nombre y apellidos', 'La CURP no corresponde con el nombre y los apellidos que escribiste. Escríbelos tal como aparecen en tu CURP.', 'Estándar de formularios (cruzada curp-nombre, sitio/formkit/js/validaciones.config.js)'),
    ('Crea tu cuenta', 'CURP', 'Error', 'Si la CURP indica segundo apellido y no se escribió', 'Tu CURP indica que tienes segundo apellido. Escríbelo tal como aparece en tu CURP.', 'Estándar de formularios (cruzada curp-segundo-apellido, sitio/formkit/js/validaciones.config.js)'),
    ('Entrar', 'Entra a tu cuenta', 'Error', 'CURP del registro de Kobo sin cuenta', 'Esta CURP tiene acreditaciones anteriores, pero todavía no tiene cuenta. Crea tu cuenta con ella para verlas.', D('config-entrar.js')),
    ('Entrar', 'Entra a tu cuenta', 'Error', 'CURP o contraseña incorrectas', 'La CURP o la contraseña no coinciden. Revisa ambas e intenta de nuevo.', D('config-entrar.js')),
    # Mis datos
    ('Mis datos', 'Tus datos', 'Etiqueta de dato', 'Siempre', 'Nombre', D('datos.js')),
    ('Mis datos', 'Tus datos', 'Etiqueta de dato', 'Siempre', 'Fecha de nacimiento', D('datos.js')),
    ('Mis datos', 'Tus datos', 'Etiqueta de dato', 'Desde el primer cotejo', 'Número de expediente', D('datos.js')),
    ('Mis datos', 'Fotografía', 'Texto', 'Fotografía confirmada', 'La Secretaría confirmó tu fotografía el DD/MM/AAAA. Si necesitas cambiarla, te toman otra en tu próxima cita.', D('datos.js')),
    ('Mis datos', 'Fotografía', 'Texto', 'Fotografía sin confirmar', 'Tu fotografía está por confirmarse en tu cita de cotejo. Mientras tanto puedes cambiarla.', D('datos.js')),
    ('Mis datos', 'Mis datos', 'Error', 'Si no se pudieron cargar los datos', 'No se pudieron cargar tus datos. Intenta de nuevo más tarde.', D('datos.js')),
    # Verificar
    ('Verificar una acreditación', 'Resultado', 'Título', 'Número que no existe', 'No encontramos esa acreditación', D('config-verificar.js')),
    ('Verificar una acreditación', 'Resultado', 'Texto', 'Número que no existe', 'Revisa que el número … esté escrito tal como aparece impreso. Si es correcto, la acreditación no está registrada en la Secretaría.', D('config-verificar.js')),
    ('Verificar una acreditación', 'Resultado', 'Estado', 'Acreditación encontrada', 'Vigente / No vigente', D('config-verificar.js')),
    ('Verificar una acreditación', 'Resultado', 'Etiqueta de dato', 'Acreditación encontrada', 'Nombre · Acreditación · Número de acreditación · Emitida el · Vence el', D('config-verificar.js')),
    ('Verificar una acreditación', 'Resultado', 'Texto', 'Acreditación vigente', 'Puede trabajar en el arbolado urbano con esta acreditación hasta la fecha de vencimiento.', D('config-verificar.js')),
    ('Verificar una acreditación', 'Resultado', 'Texto', 'Acreditación no vigente', 'Con esta acreditación no se puede trabajar en el arbolado urbano.', D('config-verificar.js')),
    # Versión
    ('Todas las pantallas', 'Pie', 'Distintivo', 'Versión de prueba', 'Versión de prueba 0.N', D('version-visible.js')),
]
# «Grupos» paso a paso (personal, de usted; decisión del 24/10/2026): lo que pinta app/grupos.js
G = 'Grupos (personal)'
dinamicos += [
    (G, 'Cabecera', 'Texto', 'Siempre', 'Hoy es DÍA N de MES de AAAA. Arme un grupo en cuatro pasos: el curso, la semana, quiénes van y confirmar.', D('grupos.js')),
    (G, 'Paso 1 · el curso', 'Texto', 'En cada curso', 'N personas en espera · N esperan 15 días o más · N ya vencidas · Falta gente para un grupo (mínimo 5) / Alcanza para N grupos nuevos de hasta 25 / Nadie espera por ahora', D('grupos.js')),
    (G, 'Paso 2 · la semana', 'Texto', 'Siempre', 'CURSO · lunes a viernes, de 9:00 a 15:00 (o «un día, de 9:00 a 15:00»). Puede haber 2 grupos al mismo tiempo, de cualquier curso; el curso empieza al menos 7 días después de confirmarlo.', D('grupos.js')),
    (G, 'Paso 2 · columna «Semana»', 'Botón', 'Semana que se puede elegir', 'Libre · 2 lugares libres / 1 lugar libre de 2 (elegida: «Elegida»). Nombre para el lector: «Elegir la semana del 2 al 6 de noviembre de 2026: 2 lugares libres»', D('grupos.js')),
    (G, 'Paso 2 · columna «Semana»', 'Estado', 'Semana que no se puede elegir', 'No disponible · Ya pasó / Ya empezó / Faltan menos de 7 días / Ya hay 2 grupos / Día inhábil / Falta más de un año', D('grupos.js')),
    (G, 'Paso 2 · calendario', 'Etiqueta', 'En los días', 'G-AAAA-NNN (folio, el primer día) · Examen (el último) · Nuevo grupo (la semana elegida) · Hoy · Inhábil', D('grupos.js')),
    (G, 'Paso 2 · semana elegida', 'Texto', 'Con una semana elegida', 'Del 2 al 6 de noviembre · Del lunes 2 de noviembre de 2026 al viernes 6 de noviembre de 2026, de 9:00 a 15:00. El viernes, repaso de conocimientos y examen final en el PILARES. (Curso de un día: «El DÍA…, de 9:00 a 15:00, con el examen ese mismo día en el PILARES».)', D('grupos.js')),
    (G, 'Paso 2 · semana elegida', 'Texto', 'Sin semana', 'Ninguna todavía · Toque «Libre» en la semana que quiere. (Grupo que ya existe: «Sin semana todavía · … o siga sin semana: el grupo la conserva como está».)', D('grupos.js')),
    (G, 'Paso 2', 'Error', 'Siguiente sin semana', 'Elija una semana libre en el calendario: toque «Libre» en la semana que quiere.', D('grupos.js')),
    (G, 'Paso 2', 'Aviso', 'Al tocar un día que no se puede', 'El DÍA no se puede: RAZÓN.', D('grupos.js')),
    (G, 'Paso 3 · quiénes van', 'Texto', 'Grupo nuevo', 'Se proponen las N personas que más esperan, primero las prioritarias y las de una misma institución de gobierno juntas. Quite o agregue a quien haga falta.', D('grupos.js')),
    (G, 'Paso 3 · quiénes van', 'Texto', 'Siempre', 'N de 25 lugares · se confirma con al menos 5 personas', D('grupos.js')),
    (G, 'Paso 3 · quiénes van', 'Plegable', 'Con más en espera que las propuestas', 'Cambiar a quiénes van (N más en espera)', D('grupos.js')),
    (G, 'Paso 3', 'Error', 'Siguiente sin nadie', 'Elija a quién va · Marque la casilla junto al nombre de al menos una persona. El grupo se confirma con al menos 5.', D('grupos.js')),
    (G, 'Paso 4 · confirmar', 'Estado', 'Antes de crearlo', 'Se crea por confirmar / Sigue por confirmar · Le falta: …; después, el área de Instructores y Capacitación le asigna la persona instructora y lo confirma. Lo puede completar en «Grupos programados» o en su ficha; mientras, las personas no ven la fecha como segura.', D('grupos.js')),
    (G, 'Paso 4 · confirmar', 'Botón', 'Grupo nuevo', 'Crear el grupo', D('grupos.js')),
    (G, 'Paso 4 · confirmar', 'Botón', 'Grupo que ya existe', 'Guardar los cambios / Agregar al grupo (confirmado)', D('grupos.js')),
    (G, 'Paso 4 · resumen', 'Valor', 'Grupo sin persona instructora', 'Persona instructora · La asigna el área de Instructores y Capacitación', D('grupos.js')),
    (G, 'Paso 1 · tipo de grupo', 'Ayuda (cambia)', 'Elegidas de una organización que no es de gobierno', 'NOMBRE no es una institución de gobierno: su gente va a un grupo combinado. Sólo las dependencias de la Ciudad, las Alcaldías y las dependencias federales tienen grupo propio.', D('grupos.js')),
    (G, 'Paso 1 · tipo de grupo', 'Ayuda (cambia)', 'Elegidas de varias organizaciones o independientes', 'El grupo es combinado: las elegidas son de varias organizaciones o independientes. Sólo las dependencias de la Ciudad, las Alcaldías y las dependencias federales tienen grupo propio, con toda su gente.', D('grupos.js')),
    (G, 'Paso 4', 'Error', 'Grupo propio de quien no es de gobierno', 'Un grupo propio es sólo de una institución de gobierno, con toda su gente. Elija «Combinado» o quite a quien no sea de esa institución.', D('grupos.js')),
    (G, 'Grupos programados', 'Aviso', 'Al crear un grupo', 'Grupo creado y confirmado: NOMBRE / Grupo creado, por confirmar: NOMBRE · FOLIO · N personas · entrada a las 09:00 · INSTRUCTOR. · Ver la ficha del grupo · Qué sigue', D('grupos.js')),
    (G, 'Grupos programados', 'Estado', 'Situación de cada grupo', 'En curso · día N de 5 / Confirmado / Por confirmar / Sin semana / Propuesta vencida / Confirmado · no inició', D('grupos.js')),
    (G, 'Grupos programados', 'Botón o enlace', 'Acciones de cada grupo', 'Completar con las N que más esperan · Elegir semana / Cambiar la semana / Elegir otra semana · Agregar personas', D('grupos.js')),
    (G, 'Grupos programados', 'Texto', 'Lo que falta', 'Semana, N personas, hora y persona instructora / Que Capacitación lo confirme / —', D('grupos.js')),
]
for p, s, t, c, x, d in dinamicos:
    quien = PERSONAL if p == G else PERSONA if p != 'Verificar una acreditación' else 'Cualquier persona (público)'
    filas.append({'pantalla': p, 'seccion': s, 'tipo': t, 'cuando': c, 'texto': x, 'quien': quien, 'donde': d, 'obs': ''})

# ── 3. Mensajes de la base al guardar o enviar ────────────────────────────────────────────────────────────
VE_PERSONA = {'anterior_no_corresponde', 'anterior_no_es_la_ultima', 'archivo_grande', 'curso_fijo', 'curso_inactivo',
              'debe_renovar', 'experiencia_insuficiente', 'falta_perfil', 'faltan_documentos', 'faltan_estudios',
              'formato_no_admitido', 'renovacion_vencida', 'requisito_ajeno', 'sin_adscripcion', 'sin_fotografia',
              'sin_poda_vigente', 'solicitud_cerrada'}
sql = (RAIZ / 'db/apd/03-apd-reglas.sql').read_text(encoding='utf8')
vistos = set()
for m in re.finditer(r"apd\.error\('([a-z_]+)',\s*(?:format\()?'((?:[^']|'')*)'", sql):
    cod, msg = m.group(1), m.group(2).replace("''", "'")
    msg = re.sub(r'%s', '…', msg)
    if (cod, msg) in vistos: continue
    vistos.add((cod, msg))
    quien = PERSONA if cod in VE_PERSONA and not (cod == 'carga_fuera_de_turno' and 'personal' in msg) else PERSONAL
    if cod == 'foto_confirmada' and msg.startswith('La Secretaría'): quien = PERSONA
    if cod == 'carga_fuera_de_turno' and msg.startswith('Sólo se cargan'): quien = PERSONA
    if cod in ('simulacion_prohibida', 'parametro_inexistente'): continue  # internos de la base de pruebas
    filas.append({'pantalla': 'Mensajes al guardar o enviar', 'seccion': 'Reglas de la base', 'tipo': 'Error',
                  'cuando': f'Regla «{cod}»', 'texto': msg, 'quien': quien,
                  'donde': 'Base de datos (db/apd/03-apd-reglas.sql)', 'obs': ''})

# ── 4. Mensajes de validación del Estándar de formularios ─────────────────────────────────────────────────
cfg = (RAIZ / 'sitio/formkit/js/validaciones.config.js').read_text(encoding='utf8')
# Mensajes de tipos de dato que este sistema no usa (RFC, coordenadas, edad, horas, código postal, mapa…): no se listan
NO_USA = re.compile(r'RFC|coordenad|latitud|longitud|[Ee]dad|hora|[Hh]ora|c[óo]digo postal|\bCP\b|colonia|mapa|[Mm]apa|cp|rango|Rango', re.I)
bloque = cfg[cfg.index('mensajes: {'):]
bloque = bloque[:bloque.index('\n    }')]
for m in re.finditer(r"^\s*([a-zA-Z]+):\s*'((?:[^'\\]|\\.)*)'", bloque, re.M):
    clave, msg = m.group(1), m.group(2).replace("\\'", "'")
    if NO_USA.search(msg) or NO_USA.search(clave) or re.match(r'(paginacion|fila|repetible|chips|cascada|compartir|sesion|hoy$)', clave): continue  # módulos que el sistema no usa
    filas.append({'pantalla': 'Mensajes del formulario (estándar)', 'seccion': 'Validación y envío', 'tipo': 'Mensaje',
                  'cuando': f'Mensaje «{clave}»', 'texto': msg, 'quien': PERSONA,
                  'donde': 'Estándar de formularios (validaciones.config.js); se cambia en el estándar o se sobrescribe en el proyecto', 'obs': ''})
for m in re.finditer(r"return '((?:[^'\\]|\\.)*[a-záéíóúñ]{3,} [^']*)'", cfg[:cfg.index('mensajes: {')]):
    msg = m.group(1).replace("\\'", "'")
    if NO_USA.search(msg): continue
    filas.append({'pantalla': 'Mensajes del formulario (estándar)', 'seccion': 'Formato de los datos', 'tipo': 'Mensaje',
                  'cuando': 'Al validar un dato con formato (correo, teléfono, CURP, etc.)', 'texto': msg, 'quien': PERSONA,
                  'donde': 'Estándar de formularios (validaciones.config.js)', 'obs': ''})

# ── 5. Observaciones de Claude sobre textos que conviene revisar ─────────────────────────────────────────────
OBS = [
    ('cargas tus documentos en PDF o en foto', 'Contradice la regla vigente: los documentos sólo se aceptan en PDF.'),
    ('Pendiente: plazo de revisión', 'Falta el plazo de la revisión documental.'),
    ('¿Olvidaste tu contraseña?', 'Falta definir quién atiende la recuperación de contraseña (pendiente).'),
    ('Fundamento: [Art. __, fracc. __]', 'Lo completa la Unidad de Transparencia; no se inventa la cita.'),
    ('Cargando', 'Texto momentáneo mientras cargan los datos.'),
    ('Termina de cargar tus documentos y envía tu solicitud.', ''),
]
for f in filas:
    for frag, obs in OBS:
        if frag in f['texto'] and obs: f['obs'] = obs
# Textos de estado que nunca se ven porque los sustituye uno más específico
for f in filas:
    if f['texto'] in ('La Secretaría emite tu acreditación.', 'Recoge tu acreditación en persona.',
                      'No alcanzaste la calificación mínima. Puedes iniciar una solicitud nueva.',
                      'No alcanzaste la asistencia mínima. Puedes iniciar una solicitud nueva.'):
        f['obs'] = 'No se muestra: lo sustituye un texto más específico de la ficha.'

# Orden de pantallas como las recorre la persona
ORDEN = ['Todas las pantallas', 'Inicio', 'Crea tu cuenta', 'Crea tu cuenta y Mis datos', 'Entrar', 'Pantallas con cuenta',
         'Mis trámites', 'Solicitud', 'Mis datos', 'Verificar una acreditación', 'Mensajes del formulario (estándar)',
         'Mensajes al guardar o enviar']
filas = [f for i, f in sorted(enumerate(filas), key=lambda e: (ORDEN.index(e[1]['pantalla']) if e[1]['pantalla'] in ORDEN else 99, e[0]))]
for i, f in enumerate(filas, 1):
    f['id'] = f'T{i:03d}'
(SAL / 'inventario-textos.json').write_text(json.dumps(filas, ensure_ascii=False, indent=1), encoding='utf8')

# ── 6. Hoja de cálculo ───────────────────────────────────────────────────────────────────────────────────────
GUINDA, GRIS, ALT, EDIT = '9D2148', '55585A', 'F7F1F3', 'FFF8E1'
fuente = 'Arial'
wb = Workbook()
guia = wb.active
guia.title = 'Cómo revisar'
hoja = wb.create_sheet('Textos')

cols = [('ID', 7), ('Pantalla', 20), ('Sección', 26), ('Tipo', 16), ('Cuándo aparece', 36), ('Texto actual', 62),
        ('Tu corrección', 52), ('Estado', 12), ('Comentario', 30), ('Observación de Claude', 34), ('Quién lo ve', 18), ('Dónde se cambia', 34)]
borde = Border(bottom=Side(style='thin', color='D9D4D6'))
for c, (nombre, ancho) in enumerate(cols, 1):
    cel = hoja.cell(row=1, column=c, value=nombre)
    cel.font = Font(name=fuente, bold=True, color='FFFFFF', size=10)
    cel.fill = PatternFill('solid', fgColor=GUINDA if nombre not in ('Tu corrección', 'Estado', 'Comentario') else 'B28E5C')
    cel.alignment = Alignment(vertical='center', wrap_text=True)
    hoja.column_dimensions[get_column_letter(c)].width = ancho
hoja.row_dimensions[1].height = 30
for r, f in enumerate(filas, 2):
    valores = [f['id'], f['pantalla'], f['seccion'], f['tipo'], f['cuando'], f['texto'], '', 'Pendiente', '', f['obs'], f['quien'], f['donde']]
    for c, v in enumerate(valores, 1):
        cel = hoja.cell(row=r, column=c, value=v)
        cel.font = Font(name=fuente, size=10, color='2B2D2E' if c not in (11, 12) else GRIS)
        cel.alignment = Alignment(vertical='top', wrap_text=True)
        cel.border = borde
        if c in (7, 8, 9):
            cel.fill = PatternFill('solid', fgColor=EDIT)
        elif r % 2 == 1:
            cel.fill = PatternFill('solid', fgColor=ALT)
    if f['obs']:
        hoja.cell(row=r, column=10).font = Font(name=fuente, size=10, color='9D2148', bold=True)
ultima = len(filas) + 1
hoja.freeze_panes = 'G2'
hoja.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{ultima}'
dv = DataValidation(type='list', formula1='"Pendiente,OK,Cambiar,Quitar,Duda"', allow_blank=False,
                    error='Elige: Pendiente, OK, Cambiar, Quitar o Duda.', errorTitle='Estado')
hoja.add_data_validation(dv)
dv.add(f'H2:H{ultima}')
hoja.conditional_formatting.add(f'H2:H{ultima}', CellIsRule(operator='equal', formula=['"OK"'], fill=PatternFill('solid', fgColor='DCEFE3'), font=Font(color='1F6B3A', bold=True)))
hoja.conditional_formatting.add(f'H2:H{ultima}', CellIsRule(operator='equal', formula=['"Cambiar"'], fill=PatternFill('solid', fgColor='FCE9C8'), font=Font(color='8A5A00', bold=True)))
hoja.conditional_formatting.add(f'H2:H{ultima}', CellIsRule(operator='equal', formula=['"Quitar"'], fill=PatternFill('solid', fgColor='F6D5DD'), font=Font(color='9D2148', bold=True)))
hoja.conditional_formatting.add(f'H2:H{ultima}', CellIsRule(operator='equal', formula=['"Duda"'], fill=PatternFill('solid', fgColor='E3E6F3'), font=Font(color='2D3E8C', bold=True)))
# Si escribe una corrección y deja «Pendiente», se resalta para recordarle cambiar el estado
hoja.conditional_formatting.add(f'G2:G{ultima}', FormulaRule(formula=[f'AND(LEN(G2)>0,H2="Pendiente")'], fill=PatternFill('solid', fgColor='FCE9C8')))

# Guía
guia.column_dimensions['A'].width = 34
guia.column_dimensions['B'].width = 14
guia.column_dimensions['C'].width = 12
guia.column_dimensions['D'].width = 12
guia.column_dimensions['E'].width = 12
guia.column_dimensions['F'].width = 12
def g(celda, valor, **kw):
    guia[celda] = valor
    guia[celda].font = Font(name=fuente, size=kw.get('size', 10), bold=kw.get('bold', False), color=kw.get('color', '2B2D2E'))
    guia[celda].alignment = Alignment(wrap_text=kw.get('wrap', False), vertical='top')
g('A1', 'Revisión de textos · Acreditaciones de arbolado urbano', size=14, bold=True, color=GUINDA)
g('A2', f'Inventario generado de la versión de prueba {(RAIZ / "VERSION").read_text().strip()} · {len(filas)} textos', color=GRIS)
pasos = [
    'Cómo revisar',
    '1. Ve a la hoja «Textos». Filtra por Pantalla, Tipo o Quién lo ve para revisar por partes (por ejemplo, primero todos los avisos).',
    '2. «Cuándo aparece» dice en qué caso ve la persona ese texto (perfil, estado de la solicitud, respuesta o error).',
    '3. Si el texto está bien, cambia Estado a «OK». Si no, escribe el texto nuevo en «Tu corrección» y pon «Cambiar».',
    '4. «Quitar» elimina el texto; «Duda» deja una pregunta en «Comentario».',
    '5. Sólo edita las columnas amarillas (Tu corrección, Estado, Comentario). Las demás vienen del sistema.',
    '6. «Observación de Claude» señala textos que contradicen una regla o tienen un pendiente.',
    '7. Devuélveme el archivo: aplico los cambios en el sitio, el catálogo de la base y los mensajes, corro las pruebas y publico la siguiente versión.',
]
for i, t in enumerate(pasos):
    guia.merge_cells(start_row=4 + i, start_column=1, end_row=4 + i, end_column=6)
    g(f'A{4 + i}', t, bold=(i == 0), wrap=True, color=GUINDA if i == 0 else '2B2D2E')
    guia.row_dimensions[4 + i].height = 15 if i == 0 else 28
fila_ej = 4 + len(pasos) + 1
g(f'A{fila_ej}', 'Ejemplo de fila revisada', bold=True, color=GUINDA)
ej = [('Texto actual', 'Llenas tu solicitud y cargas tus documentos en PDF o en foto.'),
      ('Tu corrección', 'Llenas tu solicitud y cargas tus documentos en PDF.'),
      ('Estado', 'Cambiar'), ('Comentario', 'Sólo se aceptan PDF.')]
for i, (k, v) in enumerate(ej, 1):
    g(f'A{fila_ej + i}', k, color=GRIS)
    guia.merge_cells(start_row=fila_ej + i, start_column=2, end_row=fila_ej + i, end_column=6)
    g(f'B{fila_ej + i}', v, wrap=True)
    guia[f'B{fila_ej + i}'].fill = PatternFill('solid', fgColor=EDIT)
    guia.row_dimensions[fila_ej + i].height = 28

# Avance por pantalla (fórmulas sobre la hoja Textos)
fa = fila_ej + len(ej) + 2
g(f'A{fa}', 'Avance por pantalla', bold=True, color=GUINDA)
for c, t in enumerate(['Pantalla', 'Textos', 'Pendientes', 'OK', 'Cambiar', 'Quitar o duda'], 1):
    cel = guia.cell(row=fa + 1, column=c, value=t)
    cel.font = Font(name=fuente, bold=True, color='FFFFFF', size=10)
    cel.fill = PatternFill('solid', fgColor=GUINDA)
pantallas = [p for p in ORDEN if any(f['pantalla'] == p for f in filas)]
R = f"Textos!$B$2:$B${ultima}"; E = f"Textos!$H$2:$H${ultima}"
for i, p in enumerate(pantallas + ['Total']):
    r = fa + 2 + i
    cel = guia.cell(row=r, column=1, value=p)
    cel.font = Font(name=fuente, size=10, bold=(p == 'Total'))
    if p == 'Total':
        a, b = fa + 2, fa + 1 + len(pantallas)
        for c in range(2, 7):
            L = get_column_letter(c)
            guia.cell(row=r, column=c, value=f'=SUM({L}{a}:{L}{b})').font = Font(name=fuente, size=10, bold=True)
        continue
    guia.cell(row=r, column=2, value=f'=COUNTIF({R},A{r})')
    guia.cell(row=r, column=3, value=f'=COUNTIFS({R},A{r},{E},"Pendiente")')
    guia.cell(row=r, column=4, value=f'=COUNTIFS({R},A{r},{E},"OK")')
    guia.cell(row=r, column=5, value=f'=COUNTIFS({R},A{r},{E},"Cambiar")')
    guia.cell(row=r, column=6, value=f'=COUNTIFS({R},A{r},{E},"Quitar")+COUNTIFS({R},A{r},{E},"Duda")')
    for c in range(2, 7):
        guia.cell(row=r, column=c).font = Font(name=fuente, size=10)
wb.active = 0
wb.save(SAL / 'inventario-textos.xlsx')
from collections import Counter
print(f'{len(filas)} textos → docs/textos/inventario-textos.xlsx')
print(dict(Counter(f['pantalla'] for f in filas)))
